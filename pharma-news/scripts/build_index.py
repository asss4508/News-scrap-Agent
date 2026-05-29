import json
import os
import re
import sys
from pathlib import Path

try:
    import pypdf
    def extract_pdf(path):
        text = ""
        with open(path, "rb") as f:
            reader = pypdf.PdfReader(f)
            for page in reader.pages:
                text += page.extract_text() or ""
        return text
except ImportError:
    def extract_pdf(path):
        return ""

try:
    import openpyxl
    def extract_excel(path):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"[시트: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                row_text = "\t".join(str(v) for v in row if v is not None)
                if row_text.strip():
                    lines.append(row_text)
        return "\n".join(lines)
except ImportError:
    def extract_excel(path):
        return ""

try:
    from docx import Document as DocxDocument
    def extract_word(path):
        doc = DocxDocument(path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
except ImportError:
    def extract_word(path):
        return ""

def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    elif suffix in (".xlsx", ".xls"):
        return extract_excel(path)
    elif suffix == ".docx":
        return extract_word(path)
    elif suffix in (".txt", ".md", ".csv"):
        return path.read_text(encoding="utf-8", errors="ignore")
    return ""

def split_chunks(text: str, source: str, chunk_size=800, overlap=150):
    text = re.sub(r'\n{3,}', '\n\n', text).strip()
    if not text:
        return []

    paragraphs = text.split('\n\n')
    chunks = []
    buffer = ""

    for para in paragraphs:
        para = para.strip()
        if not para:
            continue
        if len(buffer) + len(para) > chunk_size:
            if buffer:
                chunks.append({"text": buffer.strip(), "source": source})
                buffer = buffer[-overlap:] + " " + para
            else:
                # 단일 단락이 chunk_size 초과
                for i in range(0, len(para), chunk_size - overlap):
                    chunk = para[i:i + chunk_size]
                    if chunk.strip():
                        chunks.append({"text": chunk.strip(), "source": source})
                buffer = ""
        else:
            buffer = (buffer + "\n\n" + para).strip() if buffer else para

    if buffer.strip():
        chunks.append({"text": buffer.strip(), "source": source})

    return chunks

def build_index():
    base = Path(__file__).parent.parent.parent / "data"
    uploads_dir = base / "uploads"
    channels_dir = base / "channels"
    index_path = base / "index.json"

    all_chunks = []

    # 업로드된 파일 처리
    if uploads_dir.exists():
        files = [f for f in uploads_dir.rglob("*") if f.is_file() and f.name != ".gitkeep"]
        print(f"파일 처리 중: {len(files)}개")
        for f in files:
            text = extract_text(f)
            if text.strip():
                chunks = split_chunks(text, f.name)
                all_chunks.extend(chunks)
                print(f"  {f.name}: {len(chunks)}개 청크")

    # 채널 동기화 내용 처리
    if channels_dir.exists():
        channel_files = [f for f in channels_dir.glob("*.txt") if f.name != ".gitkeep"]
        print(f"채널 파일 처리 중: {len(channel_files)}개")
        for f in channel_files:
            text = f.read_text(encoding="utf-8", errors="ignore")
            if text.strip():
                chunks = split_chunks(text, f"채널:{f.stem}", chunk_size=600)
                all_chunks.extend(chunks)
                print(f"  {f.stem}: {len(chunks)}개 청크")

    with open(index_path, "w", encoding="utf-8") as fp:
        json.dump(all_chunks, fp, ensure_ascii=False, indent=2)

    print(f"\n인덱스 생성 완료: 총 {len(all_chunks)}개 청크 → data/index.json")

if __name__ == "__main__":
    build_index()
